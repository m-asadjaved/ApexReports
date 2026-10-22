"""
Script to test report generation locally on Windows without Docker.
Generates sample PDF, HTML, and XLSX outputs into the output/ folder.
"""
import sys
from pathlib import Path

# Add project root to sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook
from app.core.config import settings
from app.engine.template_processor import TemplateProcessor
from app.renderer.fast_html_renderer import FastHtmlRenderer
from app.renderer.pdf_service import PdfService

def run_sample_renders():
    settings.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    print("=" * 60)
    print("ApexReports Local Engine Test")
    print(f"Active PDF Backend: {PdfService.get_active_backend()}")
    print("=" * 60)

    # 1. Invoice Sample
    invoice_template = settings.TEMPLATES_DIR / "INVOICE_REPORT.xlsx"
    print(f"\n1. Processing Invoice template: {invoice_template.name}...")
    wb_inv = load_workbook(str(invoice_template), data_only=False)
    
    invoice_data = {
        "invoices": [
            {
                "company": {
                    "name": "Apex Enterprise Technologies Ltd."
                },
                "invoice_no": "INV-2026-9001",
                "invoice_date": "2026-10-22",
                "customer": {
                    "name": "Acme Global Corporation"
                },
                "items": [
                    {"item_code": "ORA-APX-01", "description": "Oracle APEX Custom Engine License", "qty": 5, "unit_price": 450.00},
                    {"item_code": "ORA-APX-02", "description": "Oracle Reports 6i Migration Consulting", "qty": 20, "unit_price": 125.00},
                    {"item_code": "SUP-247", "description": "24/7 Enterprise SLA Support Package", "qty": 1, "unit_price": 1500.00}
                ]
            }
        ]
    }

    processor = TemplateProcessor(wb_inv)
    populated_inv = processor.process(invoice_data)

    # Export HTML
    html_content = FastHtmlRenderer.render(populated_inv.active, title="Invoice INV-2026-9001")
    html_path = settings.OUTPUT_DIR / "sample_invoice.html"
    html_path.write_text(html_content, encoding="utf-8")
    print(f"  [OK] HTML preview generated: {html_path} ({len(html_content)} chars)")

    # Export XLSX
    xlsx_path = settings.OUTPUT_DIR / "sample_invoice.xlsx"
    populated_inv.save(str(xlsx_path))
    print(f"  [OK] Excel workbook generated: {xlsx_path} ({xlsx_path.stat().st_size} bytes)")

    # Export PDF
    pdf_bytes = PdfService.convert_to_pdf(populated_inv, filename="sample_invoice.pdf")
    pdf_path = settings.OUTPUT_DIR / "sample_invoice.pdf"
    pdf_path.write_bytes(pdf_bytes)
    print(f"  [OK] PDF document generated: {pdf_path} ({len(pdf_bytes)} bytes)")

    # 2. Master-Detail Department & Employees Sample
    dept_template = settings.TEMPLATES_DIR / "DEPT_EMPLOYEES.xlsx"
    print(f"\n2. Processing Master-Detail template: {dept_template.name}...")
    wb_dept = load_workbook(str(dept_template), data_only=False)

    dept_data = {
        "departments": [
            {
                "dept_name": "ACCOUNTING",
                "loc": "NEW YORK",
                "employees": [
                    {"empno": 7782, "ename": "CLARK", "job": "MANAGER", "sal": 2450.00},
                    {"empno": 7839, "ename": "KING", "job": "PRESIDENT", "sal": 5000.00},
                    {"empno": 7934, "ename": "MILLER", "job": "CLERK", "sal": 1300.00}
                ]
            },
            {
                "dept_name": "RESEARCH",
                "loc": "DALLAS",
                "employees": [
                    {"empno": 7369, "ename": "SMITH", "job": "CLERK", "sal": 800.00},
                    {"empno": 7566, "ename": "JONES", "job": "MANAGER", "sal": 2975.00},
                    {"empno": 7788, "ename": "SCOTT", "job": "ANALYST", "sal": 3000.00},
                    {"empno": 7876, "ename": "ADAMS", "job": "CLERK", "sal": 1100.00},
                    {"empno": 7902, "ename": "FORD", "job": "ANALYST", "sal": 3000.00}
                ]
            },
            {
                "dept_name": "SALES",
                "loc": "CHICAGO",
                "employees": [
                    {"empno": 7499, "ename": "ALLEN", "job": "SALESMAN", "sal": 1600.00},
                    {"empno": 7521, "ename": "WARD", "job": "SALESMAN", "sal": 1250.00},
                    {"empno": 7654, "ename": "MARTIN", "job": "SALESMAN", "sal": 1250.00},
                    {"empno": 7698, "ename": "BLAKE", "job": "MANAGER", "sal": 2850.00},
                    {"empno": 7844, "ename": "TURNER", "job": "SALESMAN", "sal": 1500.00}
                ]
            }
        ]
    }

    dept_processor = TemplateProcessor(wb_dept)
    populated_dept = dept_processor.process(dept_data)

    dept_pdf_bytes = PdfService.convert_to_pdf(populated_dept, filename="sample_dept_employees.pdf")
    dept_pdf_path = settings.OUTPUT_DIR / "sample_dept_employees.pdf"
    dept_pdf_path.write_bytes(dept_pdf_bytes)
    print(f"  [OK] Master-Detail PDF generated: {dept_pdf_path} ({len(dept_pdf_bytes)} bytes)")

    print("\n" + "=" * 60)
    print("All sample reports successfully generated in output/ folder!")
    print("=" * 60)

if __name__ == "__main__":
    run_sample_renders()
