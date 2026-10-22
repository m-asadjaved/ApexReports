import os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def create_invoice_template(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"

    # Styling definitions
    title_font = Font(name="Calibri", size=18, bold=True, color="1E3A8A")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True, color="1E293B")
    regular_font = Font(name="Calibri", size=11, color="334155")
    
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    subtotal_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    # Column widths
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 18

    # Row 1: Title & Barcode
    ws["A1"] = "TAX INVOICE"
    ws["A1"].font = title_font
    ws["D1"] = "{{@barcode invoice_no format=code128 width=140 height=35}}"

    # Row 2: Company Info
    ws["A2"] = "{{company.name}}"
    ws["A2"].font = bold_font
    ws["D2"] = "Invoice No: {{invoice_no}}"
    ws["D2"].font = bold_font

    # Row 3: Customer Info & Date
    ws["A3"] = "Bill To: {{customer.name}}"
    ws["A3"].font = regular_font
    ws["D3"] = "Date: {{invoice_date}}"
    ws["D3"].font = regular_font

    # Row 5: Table Headers
    headers = ["Item Code", "Description", "Quantity", "Unit Price", "Line Total"]
    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=col_idx)
        cell.value = h_text
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Row 6: Directive {{#each items}}
    ws["A6"] = "{{#each items}}"

    # Row 7: Detail Row
    ws["A7"] = "{{item_code}}"
    ws["A7"].font = regular_font
    ws["A7"].alignment = Alignment(horizontal="center")
    ws["A7"].border = thin_border

    ws["B7"] = "{{description}}"
    ws["B7"].font = regular_font
    ws["B7"].border = thin_border

    ws["C7"] = "{{qty}}"
    ws["C7"].font = regular_font
    ws["C7"].alignment = Alignment(horizontal="right")
    ws["C7"].number_format = "#,##0"
    ws["C7"].border = thin_border

    ws["D7"] = "{{unit_price}}"
    ws["D7"].font = regular_font
    ws["D7"].alignment = Alignment(horizontal="right")
    ws["D7"].number_format = "$#,##0.00"
    ws["D7"].border = thin_border

    ws["E7"] = "=C7*D7"
    ws["E7"].font = regular_font
    ws["E7"].alignment = Alignment(horizontal="right")
    ws["E7"].number_format = "$#,##0.00"
    ws["E7"].border = thin_border

    # Row 8: Directive {{/each}}
    ws["A8"] = "{{/each}}"

    # Row 9: Subtotal Row
    ws["D9"] = "Subtotal:"
    ws["D9"].font = bold_font
    ws["D9"].alignment = Alignment(horizontal="right")
    ws["E9"] = "=SUM(E7:E7)"
    ws["E9"].font = bold_font
    ws["E9"].alignment = Alignment(horizontal="right")
    ws["E9"].number_format = "$#,##0.00"
    ws["E9"].fill = subtotal_fill
    ws["E9"].border = thin_border

    # Row 10: QR Code for Verification
    ws["A10"] = "Scan to verify:"
    ws["A10"].font = bold_font
    ws["B10"] = "{{@qr 'https://verify.apex.corp/inv/' + invoice_no width=70 height=70}}"

    # Print setup
    ws.print_title_rows = "5:5"
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4

    wb.save(str(path))

def create_dept_employees_template(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Departments"

    # Fonts & styles
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    dept_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    white_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True, color="0F172A")
    regular_font = Font(name="Calibri", size=11, color="334155")

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 18

    # Row 1: Outer loop over departments
    ws["A1"] = "{{#each departments page_break=true}}"

    # Row 2: Department Header
    ws["A2"] = "DEPARTMENT: {{dept_name}} (Location: {{loc}})"
    ws["A2"].font = bold_font
    ws["A2"].fill = dept_fill

    # Row 3: Table Headers
    cols = ["Emp No", "Employee Name", "Job Title", "Salary"]
    for idx, cname in enumerate(cols, start=1):
        cell = ws.cell(row=3, column=idx)
        cell.value = cname
        cell.font = white_font
        cell.fill = header_fill

    # Row 4: Inner loop over employees
    ws["A4"] = "{{#each employees}}"

    # Row 5: Employee Detail
    ws["A5"] = "{{empno}}"
    ws["A5"].font = regular_font
    ws["B5"] = "{{ename}}"
    ws["B5"].font = regular_font
    ws["C5"] = "{{job}}"
    ws["C5"].font = regular_font
    ws["D5"] = "{{sal}}"
    ws["D5"].font = regular_font
    ws["D5"].number_format = "$#,##0.00"

    # Row 6: Inner loop close
    ws["A6"] = "{{/each}}"

    # Row 7: Department Total
    ws["C7"] = "Department Total:"
    ws["C7"].font = bold_font
    ws["D7"] = "=SUM(D5:D5)"
    ws["D7"].font = bold_font
    ws["D7"].number_format = "$#,##0.00"

    # Row 8: Outer loop close
    ws["A8"] = "{{/each}}"

    wb.save(str(path))

if __name__ == "__main__":
    templates_dir = Path("templates")
    templates_dir.mkdir(parents=True, exist_ok=True)
    create_invoice_template(templates_dir / "INVOICE_REPORT.xlsx")
    create_dept_employees_template(templates_dir / "DEPT_EMPLOYEES.xlsx")
    print("Created sample templates successfully.")
