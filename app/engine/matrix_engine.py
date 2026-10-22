import re
from copy import copy
from typing import Any, Dict, List, Optional, Tuple
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import get_column_letter, column_index_from_string
from app.engine.template_parser import ExpressionEvaluator, parse_tag_attributes, TAG_REGEX
from app.engine.cell_utils import copy_cell_style, render_cell_value

evaluator = ExpressionEvaluator()

COL_EACH_REGEX = re.compile(r"\{\{#each-col\s+([a-zA-Z0-9_\.]+)(.*?)\}\}(.*?)\{\{/each-col\}\}")

class MatrixEngine:
    """Pre-processes worksheets to horizontally expand columns for Matrix / Cross-Tab reports."""

    @classmethod
    def expand_columns(cls, ws: Worksheet, context: Dict[str, Any]) -> None:
        """Finds {{#each-col ...}} markers in columns and replicates columns horizontally."""
        # Find all cells with {{#each-col ...}}
        col_directives: List[Tuple[int, int, str, str, str]] = []

        for row in range(1, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                val = ws.cell(row=row, column=col).value
                if val and isinstance(val, str):
                    m = COL_EACH_REGEX.search(val)
                    if m:
                        list_key = m.group(1)
                        attrs_str = m.group(2)
                        inner_template = m.group(3).strip()
                        col_directives.append((row, col, list_key, attrs_str, inner_template))

        if not col_directives:
            return

        # Process directives from right to left so column insertions do not invalidate earlier column indices
        col_directives.sort(key=lambda x: x[1], reverse=True)

        for marker_row, target_col, list_key, attrs_str, inner_template in col_directives:
            items = evaluator.evaluate_expression(list_key, context)
            if not isinstance(items, (list, tuple)) or not items:
                continue

            num_items = len(items)
            if num_items <= 1:
                # 1 item: just replace the tag in place
                item_ctx = {**context, **items[0]} if isinstance(items[0], dict) else {**context, "item": items[0]}
                rendered = render_cell_value(inner_template, item_ctx)
                ws.cell(row=marker_row, column=target_col).value = rendered
                continue

            # Need to insert (num_items - 1) columns to the right of target_col
            added_cols = num_items - 1
            ws.insert_cols(target_col + 1, amount=added_cols)

            src_col_letter = get_column_letter(target_col)
            src_width = ws.column_dimensions[src_col_letter].width

            # Copy column widths and styles
            for offset in range(1, num_items):
                new_col_idx = target_col + offset
                new_col_letter = get_column_letter(new_col_idx)
                if src_width:
                    ws.column_dimensions[new_col_letter].width = src_width

                for r in range(1, ws.max_row + 1):
                    src_c = ws.cell(row=r, column=target_col)
                    dst_c = ws.cell(row=r, column=new_col_idx)
                    copy_cell_style(src_c, dst_c)
                    dst_c.value = src_c.value

            # Now populate the expanded columns for each item
            for idx, item in enumerate(items):
                curr_col_idx = target_col + idx
                item_ctx = {**context, **item} if isinstance(item, dict) else {**context, "item": item}
                item_ctx["@col_index"] = idx

                # In the header cell, set the rendered template
                header_cell = ws.cell(row=marker_row, column=curr_col_idx)
                header_cell.value = render_cell_value(inner_template, item_ctx)

                # In other cells of this column, evaluate expressions if any
                for r in range(1, ws.max_row + 1):
                    if r != marker_row:
                        c = ws.cell(row=r, column=curr_col_idx)
                        if c.value and isinstance(c.value, str) and "{{" in c.value:
                            c.value = render_cell_value(c.value, item_ctx)
