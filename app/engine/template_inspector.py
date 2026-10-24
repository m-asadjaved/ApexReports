import re
from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet


class TemplateInspector:
    """
    Analyzes an Excel report template (.xlsx) to extract metadata,
    discovered placeholders, repeating blocks, formulas, media tags,
    and inferred JSON sample payloads.
    """

    DIRECTIVE_REGEX = re.compile(r'\{\{\s*(.*?)\s*\}\}')

    @classmethod
    def inspect_file(cls, file_path: Path, template_code: Optional[str] = None) -> Dict[str, Any]:
        """Loads and inspects an Excel template file from disk."""
        wb = openpyxl.load_workbook(str(file_path), data_only=False)
        code = template_code or file_path.stem
        return cls.inspect_workbook(wb, code, filename=file_path.name)

    @classmethod
    def inspect_workbook(cls, wb: openpyxl.Workbook, template_code: str, filename: str = "") -> Dict[str, Any]:
        """Inspects an openpyxl Workbook and returns a comprehensive metadata report."""
        ws: Worksheet = wb.active
        sheet_name = ws.title

        blocks: List[Dict[str, Any]] = []
        variables: List[Dict[str, Any]] = []
        formulas: List[Dict[str, Any]] = []
        media_tags: List[Dict[str, Any]] = []
        merged_ranges: List[str] = [str(r) for r in ws.merged_cells.ranges]

        # Track block stacks to establish variable scopes
        block_stack: List[str] = []

        for row in ws.iter_rows(values_only=False):
            for cell in row:
                if cell.value is None:
                    continue

                val_str = str(cell.value).strip()

                # 1. Formulas
                if val_str.startswith("="):
                    formulas.append({
                        "cell": cell.coordinate,
                        "row": cell.row,
                        "col": cell.column_letter,
                        "formula": val_str,
                        "is_summary": any(fn in val_str.upper() for fn in ["SUM(", "COUNTA(", "COUNT(", "AVERAGE(", "MIN(", "MAX("])
                    })
                    continue

                # 2. Template Tags
                if "{{" in val_str:
                    matches = cls.DIRECTIVE_REGEX.findall(val_str)
                    for match in matches:
                        tag = match.strip()

                        # Loop Open
                        if tag.startswith("#each "):
                            parts = tag.split()
                            target = parts[1]
                            options = " ".join(parts[2:]) if len(parts) > 2 else ""
                            block_stack.append(target)
                            blocks.append({
                                "type": "loop",
                                "target": target,
                                "options": options,
                                "start_row": cell.row,
                                "start_cell": cell.coordinate
                            })

                        # Loop Close
                        elif tag == "/each":
                            closed_target = block_stack.pop() if block_stack else "unknown"
                            # Find the last open block for this target
                            for b in reversed(blocks):
                                if b.get("target") == closed_target and "end_row" not in b:
                                    b["end_row"] = cell.row
                                    b["end_cell"] = cell.coordinate
                                    break

                        # Matrix Columns
                        elif tag.startswith("#each-col "):
                            parts = tag.split()
                            target = parts[1]
                            blocks.append({
                                "type": "matrix_col",
                                "target": target,
                                "row": cell.row,
                                "cell": cell.coordinate
                            })

                        # Conditionals
                        elif tag.startswith("#if "):
                            condition = tag[4:].strip()
                            blocks.append({
                                "type": "condition",
                                "condition": condition,
                                "start_row": cell.row,
                                "start_cell": cell.coordinate
                            })

                        elif tag == "/if":
                            for b in reversed(blocks):
                                if b.get("type") == "condition" and "end_row" not in b:
                                    b["end_row"] = cell.row
                                    b["end_cell"] = cell.coordinate
                                    break

                        # Page Break
                        elif tag in ("#page-break", "page_break", "page-break"):
                            blocks.append({
                                "type": "page_break",
                                "row": cell.row,
                                "cell": cell.coordinate
                            })

                        # Repeat on page
                        elif tag in ("#repeat-on-page", "#repeat-header", "#repeat"):
                            blocks.append({
                                "type": "repeat_on_page",
                                "start_row": cell.row,
                                "start_cell": cell.coordinate
                            })

                        elif tag in ("/repeat-on-page", "/repeat-header", "/repeat"):
                            for b in reversed(blocks):
                                if b.get("type") == "repeat_on_page" and "end_row" not in b:
                                    b["end_row"] = cell.row
                                    b["end_cell"] = cell.coordinate
                                    break

                        # Media: QR / Barcode / Image
                        elif tag.startswith("@qr") or tag.startswith("@barcode") or tag.startswith("@image"):
                            parts = tag.split(None, 1)
                            mtype = parts[0][1:]  # 'qr', 'barcode', 'image'
                            params_str = parts[1] if len(parts) > 1 else ""
                            media_tags.append({
                                "cell": cell.coordinate,
                                "row": cell.row,
                                "type": mtype,
                                "params": params_str,
                                "scope": list(block_stack)
                            })

                        # Scalar Variable
                        else:
                            # Strip modifiers or filters if any
                            clean_var = tag.split()[0]
                            scope = list(block_stack)
                            variables.append({
                                "name": clean_var,
                                "cell": cell.coordinate,
                                "row": cell.row,
                                "col": cell.column_letter,
                                "scope": scope,
                                "raw_cell_value": val_str
                            })

        # Deduplicate variables while preserving scope hierarchy
        unique_vars: List[Dict[str, Any]] = []
        seen = set()
        for v in variables:
            key = (tuple(v["scope"]), v["name"])
            if key not in seen:
                seen.add(key)
                unique_vars.append(v)

        # Generate sample JSON schema
        sample_data = cls._generate_sample_data(template_code, blocks, unique_vars)

        return {
            "template_code": template_code,
            "filename": filename or f"{template_code}.xlsx",
            "sheet_name": sheet_name,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "blocks": blocks,
            "variables": unique_vars,
            "formulas": formulas,
            "media_tags": media_tags,
            "merged_ranges": merged_ranges,
            "sample_data": sample_data
        }

    @classmethod
    def _generate_sample_data(cls, template_code: str, blocks: List[Dict[str, Any]], variables: List[Dict[str, Any]]) -> Dict[str, Any]:
        """Provides rich preset data for standard templates, or auto-infers mock JSON for custom templates."""
        # Built-in presets for standard templates
        if template_code == "INVOICE_REPORT":
            return {
                "invoices": [
                    {
                        "company": {
                            "name": "Apex Enterprise Solutions Inc.",
                            "address": "100 Oracle Parkway, Redwood Shores, CA"
                        },
                        "invoice_no": "INV-2026-8801",
                        "invoice_date": "2026-10-22",
                        "customer": {
                            "name": "Oracle Corporation",
                            "email": "procurement@oracle.corp"
                        },
                        "items": [
                            {"item_code": "LIC-01", "description": "APEX Extension Enterprise License", "qty": 10, "unit_price": 450.00},
                            {"item_code": "SUP-02", "description": "24/7 Dedicated Platinum Support", "qty": 1, "unit_price": 1200.00},
                            {"item_code": "CON-03", "description": "Architecture Design Review & Tuning", "qty": 8, "unit_price": 175.00}
                        ]
                    },
                    {
                        "company": {
                            "name": "Apex Enterprise Solutions Inc.",
                            "address": "100 Oracle Parkway, Redwood Shores, CA"
                        },
                        "invoice_no": "INV-2026-8802",
                        "invoice_date": "2026-10-23",
                        "customer": {
                            "name": "Global Logistics Corp",
                            "email": "ap@globallogistics.com"
                        },
                        "items": [
                            {"item_code": "SRV-01", "description": "Cloud Integration Gateway Deployment", "qty": 1, "unit_price": 3400.00},
                            {"item_code": "TRN-04", "description": "Engineering Hands-on Workshop", "qty": 2, "unit_price": 750.00}
                        ]
                    }
                ]
            }

        if template_code == "DEPT_EMPLOYEES":
            return {
                "departments": [
                    {
                        "dept_name": "ACCOUNTING",
                        "loc": "NEW YORK",
                        "employees": [
                            {"empno": 7782, "ename": "CLARK", "job": "MANAGER", "sal": 2450.00, "manager": "KING"},
                            {"empno": 7839, "ename": "KING", "job": "PRESIDENT", "sal": 5000.00, "manager": None},
                            {"empno": 7934, "ename": "MILLER", "job": "CLERK", "sal": 1300.00, "manager": "CLARK"}
                        ]
                    },
                    {
                        "dept_name": "RESEARCH",
                        "loc": "DALLAS",
                        "employees": [
                            {"empno": 7369, "ename": "SMITH", "job": "CLERK", "sal": 800.00, "manager": "FORD"},
                            {"empno": 7566, "ename": "JONES", "job": "MANAGER", "sal": 2975.00, "manager": "KING"},
                            {"empno": 7788, "ename": "SCOTT", "job": "ANALYST", "sal": 3000.00, "manager": "JONES"},
                            {"empno": 7876, "ename": "ADAMS", "job": "CLERK", "sal": 1100.00, "manager": "SCOTT"},
                            {"empno": 7902, "ename": "FORD", "job": "ANALYST", "sal": 3000.00, "manager": "JONES"}
                        ]
                    },
                    {
                        "dept_name": "SALES",
                        "loc": "CHICAGO",
                        "employees": [
                            {"empno": 7499, "ename": "ALLEN", "job": "SALESMAN", "sal": 1600.00, "manager": "BLAKE"},
                            {"empno": 7521, "ename": "WARD", "job": "SALESMAN", "sal": 1250.00, "manager": "BLAKE"},
                            {"empno": 7654, "ename": "MARTIN", "job": "SALESMAN", "sal": 1250.00, "manager": "BLAKE"},
                            {"empno": 7698, "ename": "BLAKE", "job": "MANAGER", "sal": 2850.00, "manager": "KING"}
                        ]
                    }
                ]
            }

        # Auto-infer structure for unknown/custom templates
        payload: Dict[str, Any] = {}

        # Handle root level variables
        for v in variables:
            if not v["scope"]:
                cls._set_nested(payload, v["name"], cls._mock_value(v["name"]))

        # Handle loops
        loop_blocks = [b for b in blocks if b.get("type") == "loop"]
        if loop_blocks:
            top_loop = loop_blocks[0]["target"]
            # Find detail variables in this loop
            detail_item: Dict[str, Any] = {}
            for v in variables:
                if v["scope"] and v["scope"][0] == top_loop:
                    if len(v["scope"]) == 1:
                        cls._set_nested(detail_item, v["name"], cls._mock_value(v["name"]))

            # Handle nested inner loops if present
            if len(loop_blocks) > 1:
                inner_loop = loop_blocks[1]["target"]
                inner_item: Dict[str, Any] = {}
                for v in variables:
                    if len(v["scope"]) > 1 and v["scope"][1] == inner_loop:
                        cls._set_nested(inner_item, v["name"], cls._mock_value(v["name"]))
                detail_item[inner_loop] = [inner_item, dict(inner_item)]

            payload[top_loop] = [detail_item, dict(detail_item)]

        return payload

    @staticmethod
    def _set_nested(target_dict: Dict[str, Any], path: str, value: Any) -> None:
        """Sets a value in a nested dict using dot notation."""
        parts = path.split(".")
        curr = target_dict
        for p in parts[:-1]:
            if p not in curr or not isinstance(curr[p], dict):
                curr[p] = {}
            curr = curr[p]
        curr[parts[-1]] = value

    @staticmethod
    def _mock_value(field_name: str) -> Any:
        """Infers a sensible mock value based on the field name."""
        name_lower = field_name.lower()
        if "id" in name_lower or "no" in name_lower or "num" in name_lower:
            return 1001
        if "date" in name_lower:
            return "2026-10-22"
        if "price" in name_lower or "sal" in name_lower or "amount" in name_lower or "total" in name_lower:
            return 250.00
        if "qty" in name_lower or "count" in name_lower:
            return 5
        if "email" in name_lower:
            return "user@example.com"
        if "name" in name_lower:
            return f"Sample {field_name.capitalize()}"
        return f"Test {field_name}"
