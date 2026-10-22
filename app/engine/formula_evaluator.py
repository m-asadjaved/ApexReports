import re
from typing import Any, Dict, List, Optional, Tuple, Union
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import get_column_letter, column_index_from_string

CELL_COORD_REGEX = re.compile(r"(\$?[A-Za-z]{1,3})(\$?)(\d+)")
RANGE_REGEX = re.compile(r"(\$?[A-Za-z]{1,3}\$?\d+):(\$?[A-Za-z]{1,3}\$?\d+)")
FUNC_CALL_REGEX = re.compile(r"\b(SUM|AVERAGE|COUNT|MIN|MAX)\((.*?)\)", re.IGNORECASE)

class FormulaEvaluator:
    """Lightweight in-memory formula calculator for spreadsheet reports.
    Evaluates basic arithmetic, SUM, AVERAGE, COUNT, MIN, MAX across cells and ranges.
    """

    @classmethod
    def evaluate_worksheet(cls, ws: Worksheet) -> Dict[Tuple[int, int], Any]:
        """Calculates values for all formula cells in the worksheet.
        Returns a dictionary mapping (row, col) -> evaluated_value.
        """
        calculated_values: Dict[Tuple[int, int], Any] = {}
        max_row = ws.max_row or 1
        max_col = ws.max_column or 1

        # Iterate top-to-bottom, left-to-right so dependent cells (like totals below details) calculate in order
        for r in range(1, max_row + 1):
            for c in range(1, max_col + 1):
                cell = ws.cell(row=r, column=c)
                val = cell.value
                if val is not None and isinstance(val, str) and val.startswith("="):
                    res = cls.evaluate_formula(val, ws, calculated_values)
                    calculated_values[(r, c)] = res

        return calculated_values

    @classmethod
    def evaluate_formula(
        cls, 
        formula: str, 
        ws: Worksheet, 
        cache: Dict[Tuple[int, int], Any]
    ) -> Any:
        """Evaluates a single formula string e.g. '=C6*D6' or '=SUM(E6:E8)'."""
        if not formula or not str(formula).startswith("="):
            return formula

        expr = formula[1:].strip()

        # Step 1: Evaluate standard aggregation functions like SUM(E6:E8)
        def func_replacer(match: re.Match) -> str:
            func_name = match.group(1).upper()
            args_str = match.group(2)
            values = cls._extract_values_from_args(args_str, ws, cache)
            
            numeric_vals = [float(v) for v in values if isinstance(v, (int, float)) or cls._is_numeric(v)]
            numeric_vals = [float(v) if not isinstance(v, (int, float)) else float(v) for v in numeric_vals]

            if func_name == "SUM":
                total = sum(numeric_vals)
                return str(total)
            elif func_name == "AVERAGE":
                avg = (sum(numeric_vals) / len(numeric_vals)) if numeric_vals else 0.0
                return str(avg)
            elif func_name == "COUNT":
                return str(len(numeric_vals))
            elif func_name == "MIN":
                return str(min(numeric_vals) if numeric_vals else 0.0)
            elif func_name == "MAX":
                return str(max(numeric_vals) if numeric_vals else 0.0)
            return "0"

        expr = FUNC_CALL_REGEX.sub(func_replacer, expr)

        # Step 2: Replace individual cell references with their current/cached numeric values
        def cell_replacer(match: re.Match) -> str:
            col_str = match.group(1).replace("$", "")
            row_num = int(match.group(3))
            col_num = column_index_from_string(col_str)

            val = cache.get((row_num, col_num))
            if val is None:
                val = ws.cell(row=row_num, column=col_num).value

            if val is None or val == "":
                return "0"
            try:
                # If numeric or float
                return str(float(val))
            except (ValueError, TypeError):
                return "0"

        expr = CELL_COORD_REGEX.sub(cell_replacer, expr)

        # Step 3: Safely evaluate math expression
        try:
            # Only allow numbers, operators +, -, *, /, %, (, ), and decimals
            cleaned_expr = re.sub(r"[^0-9\+\-\*\/\%\(\)\.]", "", expr)
            if cleaned_expr:
                res = eval(cleaned_expr, {"__builtins__": None}, {})
                # Format to int if whole number
                if isinstance(res, float) and res.is_integer():
                    return int(res)
                return round(res, 4) if isinstance(res, float) else res
        except Exception:
            pass

        return formula

    @classmethod
    def _extract_values_from_args(
        cls, 
        args_str: str, 
        ws: Worksheet, 
        cache: Dict[Tuple[int, int], Any]
    ) -> List[Any]:
        """Extracts cell values from a comma-separated list of ranges or cell refs."""
        values = []
        parts = [p.strip() for p in args_str.split(",") if p.strip()]

        for p in parts:
            # Check if range like E6:E8
            m_range = RANGE_REGEX.match(p)
            if m_range:
                start_ref = m_range.group(1).replace("$", "")
                end_ref = m_range.group(2).replace("$", "")

                start_m = CELL_COORD_REGEX.match(start_ref)
                end_m = CELL_COORD_REGEX.match(end_ref)
                if start_m and end_m:
                    s_col = column_index_from_string(start_m.group(1).replace("$", ""))
                    s_row = int(start_m.group(3))
                    e_col = column_index_from_string(end_m.group(1).replace("$", ""))
                    e_row = int(end_m.group(3))

                    for r in range(min(s_row, e_row), max(s_row, e_row) + 1):
                        for c in range(min(s_col, e_col), max(s_col, e_col) + 1):
                            val = cache.get((r, c))
                            if val is None:
                                val = ws.cell(row=r, column=c).value
                            if val is not None:
                                values.append(val)
                continue

            # Single cell ref
            m_cell = CELL_COORD_REGEX.match(p)
            if m_cell:
                col_str = m_cell.group(1).replace("$", "")
                r_num = int(m_cell.group(3))
                c_num = column_index_from_string(col_str)
                val = cache.get((r_num, c_num))
                if val is None:
                    val = ws.cell(row=r_num, column=c_num).value
                if val is not None:
                    values.append(val)
                continue

            # Literal number
            if cls._is_numeric(p):
                values.append(float(p))

        return values

    @staticmethod
    def _is_numeric(val: Any) -> bool:
        try:
            float(val)
            return True
        except (ValueError, TypeError):
            return False

    @classmethod
    def format_cell_value(cls, val: Any, number_format: Optional[str]) -> str:
        """Formats a numeric value using Excel's number format mask."""
        if val is None:
            return ""

        if not isinstance(val, (int, float)):
            try:
                val = float(val)
                if val.is_integer():
                    val = int(val)
            except (ValueError, TypeError):
                return str(val)

        if not number_format or number_format.lower() in ("general", ""):
            if isinstance(val, float):
                return f"{val:,.2f}" if abs(val) >= 1000 else f"{val:.2f}"
            return f"{val:,}" if abs(val) >= 1000 else str(val)

        nf = number_format

        # Currency formats: contains $
        if "$" in nf:
            if ".00" in nf:
                return f"${val:,.2f}"
            return f"${int(round(val)):,}"

        # Percentage
        if "%" in nf:
            pct = val * 100
            if ".00" in nf:
                return f"{pct:.2f}%"
            return f"{round(pct)}%"

        # Thousands separator
        if "#,##0" in nf:
            if ".00" in nf:
                return f"{val:,.2f}"
            return f"{int(round(val)):,}"

        # Standard decimal
        if "0.00" in nf:
            return f"{val:.2f}"

        return str(val)
