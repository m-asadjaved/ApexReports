import io
import re
from typing import Any, Dict, List, Optional, Tuple, Set
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.worksheet.pagebreak import Break
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from app.engine.template_parser import ExpressionEvaluator, parse_tag_attributes, parse_media_tag, TAG_REGEX
from app.engine.cell_utils import copy_cell_style, render_cell_value
from app.engine.formula_adjuster import FormulaAdjuster
from app.engine.media_handler import MediaHandler
from app.engine.matrix_engine import MatrixEngine
from app.core.exceptions import TemplateSyntaxException

evaluator = ExpressionEvaluator()

class TemplateProcessor:
    """Core processor for Excel report templates with Oracle Reports 6i capabilities."""

    def __init__(self, workbook: Workbook):
        self.wb = workbook
        self.src_sheet: Worksheet = workbook.active
        self.target_wb = Workbook()
        self.target_sheet: Worksheet = self.target_wb.active
        self.target_sheet.title = self.src_sheet.title

        # Global row mapping: orig_row -> all target rows
        self.global_row_mapping: Dict[int, List[int]] = {}
        # Scoped row mapping: tracks the most recent rows generated in current group/iteration
        self.recent_row_mapping: Dict[int, List[int]] = {}
        # Pending formulas: (target_row, target_col, formula_str, orig_row, scoped_mapping)
        self.pending_formulas: List[Tuple[int, int, str, int, Dict[int, List[int]]]] = []
        # Page break row indices
        self.page_breaks: List[int] = []
        # Repeating header and footer target rows
        self.header_target_rows: Set[int] = set()
        self.footer_target_rows: Set[int] = set()
        self.current_page: int = 1

    def process(self, data: Dict[str, Any]) -> Workbook:
        """Processes the template using provided data dictionary and returns populated Workbook."""
        # 1. Expand Matrix columns if {{#each-col}} tags are present
        MatrixEngine.expand_columns(self.src_sheet, data)

        # 2. Copy sheet properties and column dimensions
        self._copy_sheet_setup()

        template_rows = list(self.src_sheet.iter_rows(values_only=False))
        total_rows = len(template_rows)

        total_emitted = self._execute_block(
            rows=template_rows,
            start_idx=0,
            end_idx=total_rows,
            context=data,
            target_row_start=1
        )

        self._resolve_formulas()

        for pb_row in self.page_breaks:
            self.target_sheet.row_breaks.append(Break(id=pb_row))

        self._copy_merged_cells()
        self._update_target_print_area(total_emitted - 1)
        self._update_target_conditional_formatting(total_emitted - 1)

        # Attach header/footer metadata to target sheet for downstream HTML renderers
        self.target_sheet._header_rows = set(self.header_target_rows)
        self.target_sheet._footer_rows = set(self.footer_target_rows)

        # Post-resolve @pages / total_pages across emitted cells
        total_pages = len(self.page_breaks) + 1
        max_c = self.target_sheet.max_column or 1
        for r in range(1, total_emitted):
            for c in range(1, max_c + 1):
                cell = self.target_sheet.cell(row=r, column=c)
                if cell.value and isinstance(cell.value, str):
                    if "{{@pages}}" in cell.value or "{{total_pages}}" in cell.value or "{{@total_pages}}" in cell.value:
                        cell.value = (
                            cell.value
                            .replace("{{@pages}}", str(total_pages))
                            .replace("{{total_pages}}", str(total_pages))
                            .replace("{{@total_pages}}", str(total_pages))
                        )

        return self.target_wb

    def _copy_sheet_setup(self) -> None:
        """Copies column widths, print setups, margins, and sheet formats from source to target."""
        for col_letter, col_dim in self.src_sheet.column_dimensions.items():
            self.target_sheet.column_dimensions[col_letter].width = col_dim.width
            self.target_sheet.column_dimensions[col_letter].hidden = col_dim.hidden

        try:
            if hasattr(self.src_sheet, "sheet_format"):
                if getattr(self.src_sheet.sheet_format, "defaultRowHeight", None) is not None:
                    self.target_sheet.sheet_format.defaultRowHeight = self.src_sheet.sheet_format.defaultRowHeight
                if getattr(self.src_sheet.sheet_format, "defaultColWidth", None) is not None:
                    self.target_sheet.sheet_format.defaultColWidth = self.src_sheet.sheet_format.defaultColWidth
        except Exception:
            pass

        if self.src_sheet.print_title_rows:
            self.target_sheet.print_title_rows = self.src_sheet.print_title_rows
        if self.src_sheet.print_title_cols:
            self.target_sheet.print_title_cols = self.src_sheet.print_title_cols

        try:
            if self.src_sheet.page_setup.orientation:
                self.target_sheet.page_setup.orientation = self.src_sheet.page_setup.orientation
            if self.src_sheet.page_setup.paperSize:
                self.target_sheet.page_setup.paperSize = self.src_sheet.page_setup.paperSize
        except Exception:
            pass

        try:
            if hasattr(self.src_sheet, "page_margins") and self.src_sheet.page_margins:
                self.target_sheet.page_margins.top = self.src_sheet.page_margins.top
                self.target_sheet.page_margins.bottom = self.src_sheet.page_margins.bottom
                self.target_sheet.page_margins.left = self.src_sheet.page_margins.left
                self.target_sheet.page_margins.right = self.src_sheet.page_margins.right
        except Exception:
            pass

        try:
            if getattr(self.src_sheet.page_setup, 'scale', None):
                self.target_sheet.page_setup.scale = self.src_sheet.page_setup.scale
            if getattr(self.src_sheet.page_setup, 'fitToWidth', None) is not None:
                self.target_sheet.page_setup.fitToWidth = self.src_sheet.page_setup.fitToWidth
            if getattr(self.src_sheet.page_setup, 'fitToHeight', None) is not None:
                self.target_sheet.page_setup.fitToHeight = self.src_sheet.page_setup.fitToHeight
        except Exception:
            pass

        try:
            if hasattr(self.src_sheet, 'page_margins') and self.src_sheet.page_margins:
                if getattr(self.src_sheet.page_margins, 'header', None) is not None:
                    self.target_sheet.page_margins.header = self.src_sheet.page_margins.header
                if getattr(self.src_sheet.page_margins, 'footer', None) is not None:
                    self.target_sheet.page_margins.footer = self.src_sheet.page_margins.footer
        except Exception:
            pass

        try:
            if hasattr(self.src_sheet, 'print_area') and self.src_sheet.print_area:
                self.target_sheet.print_area = self.src_sheet.print_area
        except Exception:
            pass

        try:
            if hasattr(self.src_sheet, 'print_options') and self.src_sheet.print_options:
                self.target_sheet.print_options.horizontalCentered = self.src_sheet.print_options.horizontalCentered
                self.target_sheet.print_options.verticalCentered = self.src_sheet.print_options.verticalCentered
                self.target_sheet.print_options.gridLines = self.src_sheet.print_options.gridLines
        except Exception:
            pass

        try:
            if hasattr(self.src_sheet, 'conditional_formatting'):
                for cf in self.src_sheet.conditional_formatting:
                    self.target_sheet.conditional_formatting.add(str(cf.sqref), *cf.rules)
            if hasattr(self.wb, '_differential_styles') and self.wb._differential_styles:
                self.target_wb._differential_styles = self.wb._differential_styles
        except Exception:
            pass

        try:
            if self.src_sheet.views.sheetView:
                self.target_sheet.views.sheetView[0].showGridLines = self.src_sheet.views.sheetView[0].showGridLines
        except Exception:
            pass

    def _update_target_print_area(self, final_row: int) -> None:
        """Expands target print area rows to include all emitted data rows while keeping column bounds."""
        pa = getattr(self.target_sheet, 'print_area', None)
        if not pa:
            return
        pa_str = pa[0] if isinstance(pa, (list, tuple)) else str(pa)
        if not pa_str.strip():
            return
        clean = pa_str.split('!')[-1].replace('$', '').strip()
        try:
            min_c, min_r, max_c, max_r = range_boundaries(clean)
            if min_c and max_c:
                eff_max_r = max(max_r or 1, final_row)
                start_l = get_column_letter(min_c)
                end_l = get_column_letter(max_c)
                self.target_sheet.print_area = f"{start_l}{min_r or 1}:{end_l}{eff_max_r}"
        except Exception:
            pass

    def _update_target_conditional_formatting(self, final_row: int) -> None:
        """Expands conditional formatting ranges to cover newly emitted data rows."""
        if not hasattr(self.target_sheet, "conditional_formatting") or not self.target_sheet.conditional_formatting:
            return

        for cf in list(self.target_sheet.conditional_formatting):
            new_ranges = []
            modified = False
            for rng in cf.sqref.ranges:
                if rng.max_row < final_row:
                    new_max_r = max(rng.max_row, final_row)
                    start_l = get_column_letter(rng.min_col)
                    end_l = get_column_letter(rng.max_col)
                    new_ranges.append(f"{start_l}{rng.min_row}:{end_l}{new_max_r}")
                    modified = True
                else:
                    new_ranges.append(str(rng))

            if modified:
                try:
                    rules = list(cf.rules)
                    self.target_sheet.conditional_formatting.remove(cf)
                    for nr in new_ranges:
                        self.target_sheet.conditional_formatting.add(nr, *rules)
                except Exception:
                    pass

    def _execute_block(
        self,
        rows: List[Tuple],
        start_idx: int,
        end_idx: int,
        context: Dict[str, Any],
        target_row_start: int
    ) -> int:
        """Recursively processes a block of rows."""
        target_r = target_row_start
        i = start_idx

        while i < end_idx:
            row_cells = rows[i]
            orig_row_num = i + 1

            directive = self._detect_row_directive(row_cells)

            if directive and directive["type"] == "each_start":
                close_idx = self._find_matching_close(rows, i + 1, end_idx, "each_start", "each_end")
                list_key = directive["key"]
                raw_items = evaluator.evaluate_expression(list_key, context)
                if isinstance(raw_items, (list, tuple)):
                    items = list(raw_items)
                elif isinstance(raw_items, dict):
                    items = [raw_items]
                else:
                    items = None

                is_top_level = (start_idx == 0 or target_r == 1)

                if items is None:
                    if not context:
                        # User passed empty data: {} - render 1 blank iteration so template static layout,
                        # headers, borders and table formatting are displayed instead of an empty page.
                        items = [{}]
                    elif list_key not in context and is_top_level:
                        # User passed flat single-record data without list wrapper (e.g. {"invoice_no": ...})
                        # Treat top-level context itself as the single item for this loop.
                        items = [context]
                    else:
                        items = []
                elif len(items) == 0 and not context:
                    items = [{}]

                page_break = directive.get("attrs", {}).get("page_break", False)

                # Reset recent mapping for lines inside this loop before it runs
                for r_idx in range(i + 1, close_idx):
                    self.recent_row_mapping[r_idx + 1] = []

                for item_idx, item in enumerate(items):
                    item_context = {**context, **item} if isinstance(item, dict) else {**context, "item": item, "this": item}
                    item_context["@index"] = item_idx
                    item_context["@first"] = (item_idx == 0)
                    item_context["@last"] = (item_idx == len(items) - 1)

                    if page_break and item_idx > 0:
                        self.page_breaks.append(target_r - 1)
                        self.current_page += 1

                    target_r = self._execute_block(
                        rows=rows,
                        start_idx=i + 1,
                        end_idx=close_idx,
                        context=item_context,
                        target_row_start=target_r
                    )

                i = close_idx + 1
                continue

            elif directive and directive["type"] == "if_start":
                close_idx = self._find_matching_close(rows, i + 1, end_idx, "if_start", "if_end")
                condition = directive["condition"]
                is_true = evaluator.evaluate_condition(condition, context)

                if is_true:
                    target_r = self._execute_block(
                        rows=rows,
                        start_idx=i + 1,
                        end_idx=close_idx,
                        context=context,
                        target_row_start=target_r
                    )
                i = close_idx + 1
                continue

            elif directive and directive["type"] == "header_start":
                close_idx = self._find_matching_close(rows, i + 1, end_idx, "header_start", "header_end")
                start_target_r = target_r
                target_r = self._execute_block(
                    rows=rows,
                    start_idx=i + 1,
                    end_idx=close_idx,
                    context=context,
                    target_row_start=target_r
                )
                for hr in range(start_target_r, target_r):
                    self.header_target_rows.add(hr)
                if self.header_target_rows:
                    min_h, max_h = min(self.header_target_rows), max(self.header_target_rows)
                    self.target_sheet.print_title_rows = f"${min_h}:${max_h}"
                i = close_idx + 1
                continue

            elif directive and directive["type"] == "footer_start":
                close_idx = self._find_matching_close(rows, i + 1, end_idx, "footer_start", "footer_end")
                start_target_r = target_r
                target_r = self._execute_block(
                    rows=rows,
                    start_idx=i + 1,
                    end_idx=close_idx,
                    context=context,
                    target_row_start=target_r
                )
                for fr in range(start_target_r, target_r):
                    self.footer_target_rows.add(fr)
                i = close_idx + 1
                continue

            elif directive and directive["type"] == "page_break":
                self.page_breaks.append(target_r - 1)
                self.current_page += 1
                i += 1
                continue

            elif directive and directive["type"] in ("each_end", "if_end", "header_end", "footer_end"):
                i += 1
                continue

            else:
                self._emit_row(row_cells, orig_row_num, target_r, context)
                
                # Update global and recent mappings
                if orig_row_num not in self.global_row_mapping:
                    self.global_row_mapping[orig_row_num] = []
                self.global_row_mapping[orig_row_num].append(target_r)

                if orig_row_num not in self.recent_row_mapping:
                    self.recent_row_mapping[orig_row_num] = []
                self.recent_row_mapping[orig_row_num].append(target_r)

                target_r += 1
                i += 1

        return target_r

    def _detect_row_directive(self, row_cells: Tuple) -> Optional[Dict[str, Any]]:
        """Inspects cells in a row to determine if the entire row is a directive."""
        for c in row_cells:
            val = str(c.value or "").strip()
            if not val:
                continue

            m_each = re.match(r"^\{\{#each\s+([a-zA-Z0-9_\.]+)(.*?)\}\}$", val)
            if m_each:
                key = m_each.group(1)
                attrs = parse_tag_attributes(m_each.group(2))
                return {"type": "each_start", "key": key, "attrs": attrs}

            if val == "{{/each}}":
                return {"type": "each_end"}

            m_if = re.match(r"^\{\{#if\s+(.+)\}\}$", val)
            if m_if:
                return {"type": "if_start", "condition": m_if.group(1).strip()}

            if val == "{{/if}}":
                return {"type": "if_end"}

            if val in ("{{#page-header}}", "{{#header}}"):
                return {"type": "header_start"}

            if val in ("{{/page-header}}", "{{/header}}"):
                return {"type": "header_end"}

            if val in ("{{#page-footer}}", "{{#footer}}"):
                return {"type": "footer_start"}

            if val in ("{{/page-footer}}", "{{/footer}}"):
                return {"type": "footer_end"}

            if val in ("{{#page-break}}", "{{page_break}}", "{{page-break}}"):
                return {"type": "page_break"}

        return None

    def _find_matching_close(
        self, 
        rows: List[Tuple], 
        start: int, 
        end: int, 
        open_type: str, 
        close_type: str
    ) -> int:
        depth = 1
        for idx in range(start, end):
            d = self._detect_row_directive(rows[idx])
            if d:
                if d["type"] == open_type:
                    depth += 1
                elif d["type"] == close_type:
                    depth -= 1
                    if depth == 0:
                        return idx
        raise TemplateSyntaxException(f"Unclosed '{open_type}' block starting around line {start}")

    def _emit_row(
        self, 
        src_cells: Tuple, 
        orig_row_num: int, 
        target_row_num: int, 
        context: Dict[str, Any]
    ) -> None:
        """Copies and renders a single row into the target sheet."""
        src_row_dim = self.src_sheet.row_dimensions.get(orig_row_num)
        if src_row_dim:
            if src_row_dim.height is not None:
                self.target_sheet.row_dimensions[target_row_num].height = src_row_dim.height
            if src_row_dim.hidden:
                self.target_sheet.row_dimensions[target_row_num].hidden = src_row_dim.hidden

        # Snapshot recent mapping for this row
        scoped_mapping = {k: list(v) for k, v in self.recent_row_mapping.items()}

        import datetime
        sys_context = {
            "@page": self.current_page,
            "page_no": self.current_page,
            "@date": datetime.date.today().strftime("%Y-%m-%d"),
            "current_date": datetime.date.today().strftime("%Y-%m-%d"),
            "@time": datetime.datetime.now().strftime("%H:%M:%S"),
            "current_time": datetime.datetime.now().strftime("%H:%M:%S"),
            "@timestamp": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }
        effective_context = {**sys_context, **context}

        for col_idx, src_cell in enumerate(src_cells, start=1):
            dst_cell = self.target_sheet.cell(row=target_row_num, column=col_idx)
            copy_cell_style(src_cell, dst_cell)

            raw_val = src_cell.value
            if raw_val is None:
                continue

            val_str = str(raw_val).strip()

            if val_str.startswith("="):
                self.pending_formulas.append((target_row_num, col_idx, val_str, orig_row_num, scoped_mapping))
                dst_cell.value = val_str
                continue

            val_clean = val_str.lstrip("~'").strip()
            if val_clean.startswith(("{{@image", "{{@barcode", "{{@qrcode", "{{@qr")):
                self._handle_media_tag(val_clean, dst_cell, target_row_num, col_idx, effective_context)
                continue

            dst_cell.value = render_cell_value(raw_val, effective_context)

    def _handle_media_tag(
        self, 
        tag_str: str, 
        cell: Any, 
        row: int, 
        col: int, 
        context: Dict[str, Any]
    ) -> None:
        cell_coord = f"{get_column_letter(col)}{row}"
        cell.value = ""

        m_qr = re.match(r"^\{\{@qr\s+(.+?)\}\}$", tag_str)
        if m_qr:
            raw_content = m_qr.group(1).strip()
            expr, attrs = parse_media_tag(raw_content)
            data_val = evaluator.evaluate_expression(expr, context)
            if data_val:
                req_w = attrs.get("width", 80)
                req_h = attrs.get("height", 80)
                buf = MediaHandler.generate_qr_code(str(data_val), width=req_w, height=req_h)
                MediaHandler.embed_image_in_cell(
                    self.target_sheet,
                    cell_coord,
                    buf,
                    width=req_w,
                    height=req_h
                )
            return

        m_bc = re.match(r"^\{\{@barcode\s+(.+?)\}\}$", tag_str)
        if m_bc:
            raw_content = m_bc.group(1).strip()
            expr, attrs = parse_media_tag(raw_content)
            data_val = evaluator.evaluate_expression(expr, context)
            if data_val:
                req_w = attrs.get("width", 150)
                req_h = attrs.get("height", 40)
                show_text = attrs.get("text", False)
                buf = MediaHandler.generate_barcode(
                    str(data_val), 
                    code_format=attrs.get("format", "code128"),
                    width=req_w,
                    height=req_h,
                    write_text=show_text
                )
                MediaHandler.embed_image_in_cell(
                    self.target_sheet,
                    cell_coord,
                    buf,
                    width=req_w,
                    height=req_h
                )
            return

        m_img = re.match(r"^\{\{@image\s+(.+?)\}\}$", tag_str)
        if m_img:
            raw_content = m_img.group(1).strip()
            expr, attrs = parse_media_tag(raw_content)
            src_val = evaluator.evaluate_expression(expr, context)
            if src_val:
                buf = MediaHandler.load_image_bytes(str(src_val))
                if buf:
                    MediaHandler.embed_image_in_cell(
                        self.target_sheet,
                        cell_coord,
                        buf,
                        width=attrs.get("width"),
                        height=attrs.get("height")
                    )
            return

    def _resolve_formulas(self) -> None:
        """Adapts formulas for repeating rows and expands summary ranges."""
        for target_row, target_col, formula_str, orig_row, scoped_mapping in self.pending_formulas:
            cell = self.target_sheet.cell(row=target_row, column=target_col)
            
            # Prefer scoped mapping if it exists and has the referenced row(s), else global
            effective_mapping = {**self.global_row_mapping}
            for k, v in scoped_mapping.items():
                if v:
                    effective_mapping[k] = v

            adjusted = FormulaAdjuster.adjust_formula(
                formula=formula_str,
                orig_row=orig_row,
                target_row=target_row,
                row_mapping=effective_mapping
            )
            cell.value = adjusted

    def _copy_merged_cells(self) -> None:
        for rng in self.src_sheet.merged_cells.ranges:
            min_r, min_c, max_r, max_c = rng.min_row, rng.min_col, rng.max_row, rng.max_col
            if min_r in self.global_row_mapping and max_r in self.global_row_mapping:
                mapped_start = self.global_row_mapping[min_r]
                mapped_end = self.global_row_mapping[max_r]

                # Case 1: Single row merge across columns (e.g. A2:D2)
                if min_r == max_r:
                    for target_r in mapped_start:
                        try:
                            self.target_sheet.merge_cells(
                                start_row=target_r,
                                start_column=min_c,
                                end_row=target_r,
                                end_column=max_c
                            )
                        except Exception:
                            pass
                # Case 2: Multi-row block merge
                else:
                    if len(mapped_start) == len(mapped_end):
                        for idx in range(len(mapped_start)):
                            try:
                                self.target_sheet.merge_cells(
                                    start_row=mapped_start[idx],
                                    start_column=min_c,
                                    end_row=mapped_end[idx],
                                    end_column=max_c
                                )
                            except Exception:
                                pass
