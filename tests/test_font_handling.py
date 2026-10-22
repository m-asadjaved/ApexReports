import openpyxl
from openpyxl.styles import Font
from fastapi.testclient import TestClient
from main import app
from app.renderer.fast_html_renderer import FastHtmlRenderer

client = TestClient(app)

def test_excel_font_preserved_in_html():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Calibri Body"
    ws["A1"].font = Font(name="Calibri")
    
    ws["A2"] = "Times New Roman Heading"
    ws["A2"].font = Font(name="Times New Roman", size=16, bold=True)
    
    html = FastHtmlRenderer.render(ws)
    # Primary font in body
    assert "Calibri" in html
    # Distinct cell font on A2
    assert "Times New Roman" in html

def test_custom_font_family_override_parameter():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Override Test"
    
    html = FastHtmlRenderer.render(ws, font_family="Arial")
    assert "'Arial', 'Liberation Sans'" in html

def test_api_render_with_font_family_override():
    payload = {
        "template_code": "INVOICE_REPORT",
        "format": "html",
        "response_mode": "binary",
        "font_family": "Tahoma",
        "data": {
            "invoices": [
                {
                    "company": {"name": "Font Test Corp"},
                    "invoice_no": "INV-FNT-1",
                    "invoice_date": "2026-10-23",
                    "customer": {"name": "Customer"},
                    "items": []
                }
            ]
        }
    }
    response = client.post("/api/v1/reports/render", json=payload)
    assert response.status_code == 200
    assert "Tahoma" in response.text
