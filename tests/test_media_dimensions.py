from openpyxl import Workbook
from app.engine.template_parser import ExpressionEvaluator, parse_media_tag
from app.engine.template_processor import TemplateProcessor
from app.renderer.fast_html_renderer import FastHtmlRenderer

def test_parse_media_tag_with_spaces_and_quotes():
    tag_str = "'https://verify.corp.com/id=' + invoice_id width=140 height=140"
    expr, attrs = parse_media_tag(tag_str)
    assert expr == "'https://verify.corp.com/id=' + invoice_id"
    assert attrs["width"] == 140
    assert attrs["height"] == 140

def test_media_custom_width_height():
    wb = Workbook()
    ws = wb.active

    ws["A1"] = "{{@qr 'ABC-123' width=180 height=180}}"
    ws["B1"] = "{{@barcode 'SKU-999' format=code128 width=260 height=55}}"

    processor = TemplateProcessor(wb)
    out_wb = processor.process({})
    out_ws = out_wb.active

    images = out_ws._images
    assert len(images) == 2

    qr_img = images[0]
    bc_img = images[1]

    assert qr_img.width == 180
    assert qr_img.height == 180

    assert bc_img.width == 260
    assert bc_img.height == 55

    html = FastHtmlRenderer.render(out_ws)
    assert "width: 180px; height: 180px;" in html
    assert "width: 260px; height: 55px;" in html

def test_qr_string_concatenation_with_number_and_string():
    ev = ExpressionEvaluator()
    expr, attrs = parse_media_tag("'https://verify.apex.corp/inv/' + invoice_no width=90 height=90")
    
    # 1. Test when invoice_no is an INTEGER (common in Oracle DB e.g. 10024)
    res_int = ev.evaluate_expression(expr, {"invoice_no": 10024})
    assert res_int == "https://verify.apex.corp/inv/10024"

    # 2. Test when invoice_no is a STRING
    res_str = ev.evaluate_expression(expr, {"invoice_no": "INV-2026-9001"})
    assert res_str == "https://verify.apex.corp/inv/INV-2026-9001"

    # 3. Test embedded token syntax
    res_token = ev.evaluate_expression("https://verify.apex.corp/inv/{{invoice_no}}", {"invoice_no": "ABC"})
    assert res_token == "https://verify.apex.corp/inv/ABC"

def test_barcode_with_text():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "{{@barcode 'INV-9901' format=code128 text=true width=180 height=50}}"

    processor = TemplateProcessor(wb)
    out_wb = processor.process({})
    assert len(out_wb.active._images) == 1
    assert out_wb.active._images[0].width == 180
    assert out_wb.active._images[0].height == 50
