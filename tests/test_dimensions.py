import openpyxl
from openpyxl.styles import Alignment
from app.engine.template_processor import TemplateProcessor
from app.renderer.fast_html_renderer import FastHtmlRenderer
from app.renderer.browser_pdf_renderer import BrowserPdfRenderer


def test_column_widths_to_colgroup():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TestDimensions"

    # Set column widths: Col A = 16.0, Col B = 34.0, Col C = 10.0
    ws.column_dimensions["A"].width = 16.0
    ws.column_dimensions["B"].width = 34.0
    ws.column_dimensions["C"].width = 10.0

    ws["A1"] = "Col A"
    ws["B1"] = "Col B"
    ws["C1"] = "Col C"

    # 1. Test fit_to_page=True (Default): All columns fit to 100% table width with exact proportions
    html = FastHtmlRenderer.render(ws, "Dimensions Test", fit_to_page=True)

    # Formula: round(w * 7 + 5)
    # A (16.0): 117px -> 26.90%
    # B (34.0): 243px -> 55.86%
    # C (10.0): 75px  -> 17.24% (100.0 - 26.90 - 55.86)
    assert "<colgroup>" in html
    assert '<col style="width: 26.90%;" />' in html
    assert '<col style="width: 55.86%;" />' in html
    assert '<col style="width: 17.24%;" />' in html
    assert "table-layout: fixed;" in html
    assert 'style="width: 100%;"' in html

    # 2. Test fit_to_page=False: Explicit pixel widths
    html_px = FastHtmlRenderer.render(ws, "Dimensions Test", fit_to_page=False)
    assert '<col style="width: 117px;" />' in html_px
    assert '<col style="width: 243px;" />' in html_px
    assert '<col style="width: 75px;" />' in html_px
    assert "min-width: 435px;" in html_px


def test_row_heights_and_default_height():
    wb = openpyxl.Workbook()
    ws = wb.active

    ws.sheet_format.defaultRowHeight = 18.0
    ws.row_dimensions[1].height = 36.5
    ws.row_dimensions[2].height = 22.0

    ws["A1"] = "Header"
    ws["A2"] = "Subheader"
    ws["A3"] = "Normal Row"

    html = FastHtmlRenderer.render(ws, "Row Heights Test")

    # Verify default row height in CSS
    assert "table.report-table tr { height: 18.0pt; }" in html
    # Verify explicit row heights inline on <tr>
    assert '<tr style="height: 36.5pt;">' in html
    assert '<tr style="height: 22.0pt;">' in html
    # Row 3 uses default height (no custom inline height on <tr>)
    assert '<tr>' in html


def test_hidden_rows_and_hidden_columns():
    wb = openpyxl.Workbook()
    ws = wb.active

    ws.column_dimensions["A"].width = 20.0
    ws.column_dimensions["B"].hidden = True
    ws.row_dimensions[2].hidden = True

    ws["A1"] = "Visible 1"
    ws["B1"] = "Hidden Col"
    ws["A2"] = "Hidden Row"
    ws["B2"] = "Hidden Both"

    html = FastHtmlRenderer.render(ws, "Hidden Test")

    # Hidden column in colgroup
    assert '<col style="display: none; width: 0px;" />' in html
    # Cell in hidden column
    assert "display: none;" in html
    # Hidden row
    assert '<tr style="display: none;">' in html


def test_cloned_rows_preserve_height():
    wb = openpyxl.Workbook()
    ws = wb.active

    ws["A1"] = "{{#each items}}"
    ws["A2"] = "{{name}}"
    ws["A3"] = "{{/each}}"

    # Give row 2 in template explicit height 28.5pt
    ws.row_dimensions[2].height = 28.5

    proc = TemplateProcessor(wb)
    out_wb = proc.process({
        "items": [
            {"name": "Item 1"},
            {"name": "Item 2"},
            {"name": "Item 3"}
        ]
    })
    out_ws = out_wb.active

    assert out_ws.row_dimensions[1].height == 28.5
    assert out_ws.row_dimensions[2].height == 28.5
    assert out_ws.row_dimensions[3].height == 28.5

    html = FastHtmlRenderer.render(out_ws)
    assert html.count('style="height: 28.5pt;"') == 3


def test_cell_vertical_alignment_and_wrapping():
    wb = openpyxl.Workbook()
    ws = wb.active

    ws["A1"] = "Centered Text"
    ws["A1"].alignment = Alignment(vertical="center", horizontal="center")

    ws["B1"] = "Line 1\nLine 2"
    ws["B1"].alignment = Alignment(wrap_text=True)

    html = FastHtmlRenderer.render(ws)

    # In CSS, Excel 'center' vertical-alignment must map to 'middle'
    assert "vertical-align: middle" in html
    assert "text-align: center" in html
    # Wrap text
    assert "white-space: pre-wrap" in html
    # Newline in cell content converted to <br/>
    assert "Line 1<br/>Line 2" in html


def test_pdf_generation_with_dimensions():
    wb = openpyxl.Workbook()
    ws = wb.active

    ws.column_dimensions["A"].width = 25.0
    ws.column_dimensions["B"].width = 30.0
    ws.row_dimensions[1].height = 40.0

    ws["A1"] = "Report Title"
    ws["B1"] = "Status"
    ws["A2"] = "Widget 1"
    ws["B2"] = "Active"

    pdf_bytes = BrowserPdfRenderer.convert_worksheet_to_pdf(ws, "Test Dimension PDF")
    assert pdf_bytes.startswith(b"%PDF")
    assert len(pdf_bytes) > 1000


def test_print_area_exclusion_of_outer_columns():
    """Verify that columns outside the Excel print_area are excluded from the HTML table."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Col A"
    ws["B1"] = "Col B"
    ws["C1"] = "Col C (Outside)"
    ws["D1"] = "Col D (Outside)"

    ws.column_dimensions["A"].width = 20.0
    ws.column_dimensions["B"].width = 25.0
    ws.column_dimensions["C"].width = 30.0
    ws.column_dimensions["D"].width = 15.0

    # Explicit print area: only A1:B1
    ws.print_area = "A1:B1"

    html = FastHtmlRenderer.render(ws, "Print Area Test")

    # Only 2 cols should be in colgroup
    assert html.count("<col style=") == 2
    assert "Col A" in html
    assert "Col B" in html
    assert "Col C (Outside)" not in html
    assert "Col D (Outside)" not in html


def test_horizontal_centering_and_margin_balancing():
    """Verify that all reports have margin-left: auto and margin-right: auto and balanced @page margins."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Header"
    ws.column_dimensions["A"].width = 20.0

    html = FastHtmlRenderer.render(ws, "Center Test")

    # Table must be horizontally centered
    assert "margin-left: auto !important; margin-right: auto !important;" in html
    assert "display: flex; flex-direction: column; align-items: center;" in html

    # @page rule must exist with balanced horizontal margin
    assert "@page {" in html
    assert "size: a4 portrait;" in html


def test_template_processor_updates_print_area_on_expansion():
    """Verify that TemplateProcessor preserves print area and expands rows to cover emitted items."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "{{#each items}}"
    ws["A2"] = "{{name}}"
    ws["B2"] = "{{val}}"
    ws["C2"] = "Scratch Column"
    ws["A3"] = "{{/each}}"

    # Print area is columns A:B
    ws.print_area = "A1:B3"

    proc = TemplateProcessor(wb)
    out_wb = proc.process({
        "items": [
            {"name": "Item 1", "val": 10},
            {"name": "Item 2", "val": 20},
            {"name": "Item 3", "val": 30},
            {"name": "Item 4", "val": 40},
            {"name": "Item 5", "val": 50},
        ]
    })
    out_ws = out_wb.active

    # Rows expanded from 3 to 5
    assert "A$1:$B$5" in out_ws.print_area or "A1:B5" in out_ws.print_area.replace("$", "")
    html = FastHtmlRenderer.render(out_ws)
    # Column C must be excluded
    assert "Scratch Column" not in html
    assert html.count("<col style=") == 2


def test_merged_cells_spans_and_borders():
    """Verify merged cell colspan/rowspan accuracy and outer boundary border preservation."""
    from openpyxl.styles import Border, Side
    from app.engine.compiled_template import TemplateCompiler

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MergeTest"

    # Merge A1:D1 (1 row x 4 cols: colspan=4, rowspan=1)
    ws.merge_cells("A1:D1")
    ws["A1"] = "Horizontal Merge Title"

    # Merge A2:B4 (3 rows x 2 cols: colspan=2, rowspan=3)
    ws.merge_cells("A2:B4")
    ws["A2"] = "Block Merge"

    # Put a border on the outer edge (e.g. cell D1 right border, cell B4 bottom border)
    thin_border = Border(
        right=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )
    ws["D1"].border = thin_border
    ws["B4"].border = thin_border

    # Test 1: FastHtmlRenderer
    html = FastHtmlRenderer.render(ws, "Merge Test")
    assert '<td colspan="4"' in html
    assert 'Horizontal Merge Title' in html
    assert 'colspan="2"' in html
    assert 'rowspan="3"' in html
    assert 'Block Merge' in html
    # Boundary borders preserved
    assert "border-right: 1px solid #000000" in html
    assert "border-bottom: 1px solid #000000" in html

    # Test 2: TemplateCompiler & CompiledRenderer
    from app.engine.compiled_template import CompiledRenderer
    compiled = TemplateCompiler.compile(wb, template_code="TEST_MERGE")
    compiled_html = CompiledRenderer.render(compiled, {})
    assert '<td colspan="4"' in compiled_html
    assert 'Horizontal Merge Title' in compiled_html
    assert 'colspan="2"' in compiled_html
    assert 'rowspan="3"' in compiled_html
    assert 'Block Merge' in compiled_html
    assert "border-right: 1px solid #000000" in compiled_html
    assert "border-bottom: 1px solid #000000" in compiled_html

