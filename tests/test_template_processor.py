import io
import pytest
from openpyxl import Workbook
from app.engine.template_processor import TemplateProcessor

def test_template_processor_basic_and_formulas():
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"

    # Row 1: Header
    ws["A1"] = "Company: {{company}}"
    # Row 2: Directive {{#each items}}
    ws["A2"] = "{{#each items}}"
    # Row 3: Template row
    ws["A3"] = "{{name}}"
    ws["B3"] = "{{qty}}"
    ws["C3"] = "{{price}}"
    ws["D3"] = "=B3*C3"
    # Row 4: Directive {{/each}}
    ws["A4"] = "{{/each}}"
    # Row 5: Summary row
    ws["A5"] = "Grand Total:"
    ws["D5"] = "=SUM(D3:D3)"
    # Row 6: Conditional block
    ws["A6"] = "{{#if has_discount}}"
    ws["A7"] = "Special Discount: {{discount}}%"
    ws["A8"] = "{{/if}}"

    data = {
        "company": "Apex Global Solutions",
        "has_discount": True,
        "discount": 15,
        "items": [
            {"name": "Laptop", "qty": 2, "price": 1200},
            {"name": "Mouse", "qty": 5, "price": 25},
            {"name": "Monitor", "qty": 1, "price": 300},
        ]
    }

    processor = TemplateProcessor(wb)
    out_wb = processor.process(data)
    out_ws = out_wb.active

    # Check header
    assert out_ws["A1"].value == "Company: Apex Global Solutions"

    # Check repeated rows (3 items: target rows 2, 3, 4)
    assert out_ws["A2"].value == "Laptop"
    assert out_ws["B2"].value == 2
    assert out_ws["C2"].value == 1200
    assert out_ws["D2"].value == "=B2*C2"

    assert out_ws["A3"].value == "Mouse"
    assert out_ws["B3"].value == 5
    assert out_ws["C3"].value == 25
    assert out_ws["D3"].value == "=B3*C3"

    assert out_ws["A4"].value == "Monitor"
    assert out_ws["B4"].value == 1
    assert out_ws["C4"].value == 300
    assert out_ws["D4"].value == "=B4*C4"

    # Check Grand Total row (target row 5)
    assert out_ws["A5"].value == "Grand Total:"
    assert out_ws["D5"].value == "=SUM(D2:D4)"

    # Check conditional block rendered (target row 6)
    assert out_ws["A6"].value == "Special Discount: 15%"

def test_template_processor_nested_master_detail():
    wb = Workbook()
    ws = wb.active
    ws.title = "Departments"

    # Row 1: {{#each departments}}
    ws["A1"] = "{{#each departments}}"
    # Row 2: Dept Header
    ws["A2"] = "Dept: {{dept_name}}"
    # Row 3: {{#each employees}}
    ws["A3"] = "{{#each employees}}"
    # Row 4: Emp detail
    ws["A4"] = "{{ename}}"
    ws["B4"] = "{{sal}}"
    # Row 5: {{/each}}
    ws["A5"] = "{{/each}}"
    # Row 6: Dept Total
    ws["A6"] = "Dept Total:"
    ws["B6"] = "=SUM(B4:B4)"
    # Row 7: {{/each}}
    ws["A7"] = "{{/each}}"

    data = {
        "departments": [
            {
                "dept_name": "ACCOUNTING",
                "employees": [
                    {"ename": "CLARK", "sal": 2450},
                    {"ename": "KING", "sal": 5000}
                ]
            },
            {
                "dept_name": "RESEARCH",
                "employees": [
                    {"ename": "SMITH", "sal": 800},
                    {"ename": "FORD", "sal": 3000},
                    {"ename": "SCOTT", "sal": 3000}
                ]
            }
        ]
    }

    processor = TemplateProcessor(wb)
    out_wb = processor.process(data)
    out_ws = out_wb.active

    # Accounting Dept:
    # Row 1: Dept: ACCOUNTING
    # Row 2: CLARK, 2450
    # Row 3: KING, 5000
    # Row 4: Dept Total:, =SUM(B2:B3)
    assert out_ws["A1"].value == "Dept: ACCOUNTING"
    assert out_ws["A2"].value == "CLARK"
    assert out_ws["B2"].value == 2450
    assert out_ws["A3"].value == "KING"
    assert out_ws["B3"].value == 5000
    assert out_ws["A4"].value == "Dept Total:"
    assert out_ws["B4"].value == "=SUM(B2:B3)"

    # Research Dept:
    # Row 5: Dept: RESEARCH
    # Row 6: SMITH, 800
    # Row 7: FORD, 3000
    # Row 8: SCOTT, 3000
    # Row 9: Dept Total:, =SUM(B6:B8)
    assert out_ws["A5"].value == "Dept: RESEARCH"
    assert out_ws["A6"].value == "SMITH"
    assert out_ws["A7"].value == "FORD"
    assert out_ws["A8"].value == "SCOTT"
    assert out_ws["A9"].value == "Dept Total:"
    assert out_ws["B9"].value == "=SUM(B6:B8)"

def test_template_processor_counta():
    wb = Workbook()
    ws = wb.active
    ws.title = "EmpCount"

    ws["A1"] = "{{#each employees}}"
    ws["A2"] = "{{ename}}"
    ws["A3"] = "{{/each}}"
    ws["A4"] = "Total Staff:"
    ws["B4"] = "=COUNTA(A2:A2)"

    data = {
        "employees": [
            {"ename": "SMITH"},
            {"ename": "ALLEN"},
            {"ename": "WARD"}
        ]
    }

    processor = TemplateProcessor(wb)
    out_wb = processor.process(data)
    out_ws = out_wb.active

    # Target rows 1, 2, 3 are employees
    assert out_ws["A1"].value == "SMITH"
    assert out_ws["A2"].value == "ALLEN"
    assert out_ws["A3"].value == "WARD"
    
    # Target row 4: Total Staff
    assert out_ws["A4"].value == "Total Staff:"
    # Formula range was auto-expanded to A1:A3
    assert out_ws["B4"].value == "=COUNTA(A1:A3)"

    # HTML rendering evaluates COUNTA to 3
    from app.renderer.fast_html_renderer import FastHtmlRenderer
    html = FastHtmlRenderer.render(out_ws)
    assert ">3</td>" in html


def test_template_processor_empty_data_preserves_static_cells():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "{{#each invoices}}"
    ws["A2"] = "SALES TAX INVOICE"
    ws["B2"] = "Invoice: {{inv_no}}"
    ws["A3"] = "{{#each items}}"
    ws["A4"] = "{{item_name}}"
    ws["B4"] = "{{qty}}"
    ws["A5"] = "{{/each}}"
    ws["A6"] = "Total:"
    ws["A7"] = "{{/each}}"

    # When caller passes completely empty data: {}
    processor = TemplateProcessor(wb)
    out_wb = processor.process({})
    out_ws = out_wb.active

    # Static cells must NOT be wiped out!
    assert out_ws.max_row >= 2
    assert out_ws["A1"].value == "SALES TAX INVOICE"
    assert out_ws["B1"].value == "Invoice: "
    assert any(out_ws.cell(r, 1).value == "Total:" for r in range(1, out_ws.max_row + 1))


def test_template_processor_flat_single_record():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "{{#each invoices}}"
    ws["A2"] = "Invoice: {{inv_no}}"
    ws["A3"] = "{{#each items}}"
    ws["A4"] = "{{item_name}}"
    ws["A5"] = "{{/each}}"
    ws["A6"] = "{{/each}}"

    # When caller passes a single flat record without "invoices" array wrapper
    data = {
        "inv_no": "INV-777",
        "items": [{"item_name": "Item A"}, {"item_name": "Item B"}]
    }

    processor = TemplateProcessor(wb)
    out_wb = processor.process(data)
    out_ws = out_wb.active

    assert out_ws["A1"].value == "Invoice: INV-777"
    assert out_ws["A2"].value == "Item A"
    assert out_ws["A3"].value == "Item B"
