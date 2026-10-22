from openpyxl import Workbook
from app.engine.template_processor import TemplateProcessor

def test_matrix_crosstab_column_expansion():
    wb = Workbook()
    ws = wb.active
    ws.title = "SalesMatrix"

    # Row 1: Headers
    ws["A1"] = "Product"
    ws["B1"] = "{{#each-col quarters}} {{name}} {{/each-col}}"
    ws["C1"] = "Total"

    # Row 2: Repeater row for products
    ws["A2"] = "{{#each products}}"
    ws["A3"] = "{{pname}}"
    ws["C3"] = "TotalSales"
    ws["A4"] = "{{/each}}"

    data = {
        "quarters": [
            {"name": "Q1 2026"},
            {"name": "Q2 2026"},
            {"name": "Q3 2026"},
            {"name": "Q4 2026"}
        ],
        "products": [
            {"pname": "Enterprise ERP"},
            {"pname": "Cloud Storage"}
        ]
    }

    processor = TemplateProcessor(wb)
    out_wb = processor.process(data)
    out_ws = out_wb.active

    # Check header expansion
    # Column A: Product
    # Column B: Q1 2026
    # Column C: Q2 2026
    # Column D: Q3 2026
    # Column E: Q4 2026
    # Column F: Total (shifted by 3 columns)
    assert out_ws["A1"].value == "Product"
    assert out_ws["B1"].value == "Q1 2026"
    assert out_ws["C1"].value == "Q2 2026"
    assert out_ws["D1"].value == "Q3 2026"
    assert out_ws["E1"].value == "Q4 2026"
    assert out_ws["F1"].value == "Total"

    # Check product rows
    assert out_ws["A2"].value == "Enterprise ERP"
    assert out_ws["A3"].value == "Cloud Storage"
