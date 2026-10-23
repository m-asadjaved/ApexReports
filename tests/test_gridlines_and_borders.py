import openpyxl
from openpyxl.styles import Border, Side
from app.renderer.fast_html_renderer import FastHtmlRenderer

def test_default_rendering_has_no_gridlines():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Invoice Title"
    ws["A2"] = "Customer Name"

    html_output = FastHtmlRenderer.render(ws)
    # Ensure the table does not have the show-gridlines class
    assert 'class="report-table"' in html_output
    assert 'class="report-table show-gridlines"' not in html_output
    # Ensure cells do not have an artificial inline border in body
    body_content = html_output.split('<body>')[1]
    assert 'style="border' not in body_content
    assert 'border-top:' not in body_content

def test_explicit_cell_borders_are_rendered():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Header"
    
    thin_side = Side(style="thin", color="000000")
    ws["A1"].border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
    
    html_output = FastHtmlRenderer.render(ws)
    assert 'border-top: 1px solid #000000' in html_output
    assert 'border-bottom: 1px solid #000000' in html_output
    assert 'border-left: 1px solid #000000' in html_output
    assert 'border-right: 1px solid #000000' in html_output

def test_print_options_gridlines_enabled():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Data"
    ws.print_options.gridLines = True
    
    html_output = FastHtmlRenderer.render(ws)
    assert 'class="report-table show-gridlines"' in html_output

def test_auto_border_color_and_directive_footer_border_inheritance():
    from app.engine.template_processor import TemplateProcessor
    from app.engine.compiled_template import TemplateCompiler, CompiledRenderer
    
    wb = openpyxl.Workbook()
    ws = wb.active
    
    # Header row
    ws["A1"] = "Item"
    ws["B1"] = "Amount"
    
    # Body row
    ws["A2"] = "{{#each items}}"
    ws["A3"] = "{{name}}"
    ws["B3"] = "{{amount}}"
    ws["B3"].number_format = "#,##0"
    ws["A4"] = "{{/each}}"
    
    # Footer with border placed on footer_start directive row
    ws["A5"] = "{{#page-footer}}"
    thick_side = Side(style="thick")  # auto color (no explicit color)
    ws["A5"].border = Border(bottom=thick_side)
    ws["B5"].border = Border(bottom=thick_side)
    ws["A6"] = "Total Footer"
    ws["A7"] = "{{/page-footer}}"
    
    data = {"items": [{"name": "Widget", "amount": 1000}]}
    
    # 1. Full processor mode
    proc = TemplateProcessor(wb)
    out_wb = proc.process(data)
    full_html = FastHtmlRenderer.render(out_wb.active)
    
    assert "border-top: 3px solid #000000" in full_html
    assert "1,000" in full_html
    
    # 2. Fast compiled mode
    ast = TemplateCompiler.compile(wb, "TEST_BORDER_INHERIT")
    fast_html = CompiledRenderer.render(ast, data)
    
    assert "border-top: 3px solid #000000" in fast_html
    assert "1,000" in fast_html

