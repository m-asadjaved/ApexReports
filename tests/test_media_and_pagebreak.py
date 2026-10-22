from openpyxl import Workbook
from app.engine.template_processor import TemplateProcessor

def test_template_processor_media_and_page_break():
    wb = Workbook()
    ws = wb.active
    ws.title = "Badge"

    ws["A1"] = "Employee: {{name}}"
    ws["B1"] = "{{@qr 'https://company.internal/emp/' + id width=60 height=60}}"
    ws["C1"] = "{{@barcode id format=code128 width=120 height=30}}"
    ws["A2"] = "{{#page-break}}"
    ws["A3"] = "Next Page Notice"

    data = {
        "name": "Jane Doe",
        "id": "EMP-9901"
    }

    processor = TemplateProcessor(wb)
    out_wb = processor.process(data)
    out_ws = out_wb.active

    assert out_ws["A1"].value == "Employee: Jane Doe"
    # QR code and barcode should clear text and embed images
    assert out_ws["B1"].value == ""
    assert out_ws["C1"].value == ""
    # Should have 2 embedded images
    assert len(out_ws._images) == 2

    # Check page break
    assert len(processor.page_breaks) == 1
    assert processor.page_breaks[0] == 1  # after row 1
