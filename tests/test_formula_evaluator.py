from openpyxl import Workbook
from app.engine.formula_evaluator import FormulaEvaluator

def test_counta_evaluation():
    wb = Workbook()
    ws = wb.active

    # Rows 9, 10, 11 contain employee names
    ws["A9"] = "CLARK"
    ws["A10"] = "KING"
    ws["A11"] = "MILLER"

    # Row 12 has COUNTA formula
    ws["A12"] = "=COUNTA(A9:A11)"

    cache = FormulaEvaluator.evaluate_worksheet(ws)
    assert cache[(12, 1)] == 3

def test_count_and_subtotal():
    wb = Workbook()
    ws = wb.active

    ws["B1"] = 10
    ws["B2"] = 20
    ws["B3"] = 30

    ws["B4"] = "=SUM(B1:B3)"
    ws["B5"] = "=AVERAGE(B1:B3)"
    ws["B6"] = "=COUNT(B1:B3)"
    ws["B7"] = "=ROUND(B5, 0)"

    cache = FormulaEvaluator.evaluate_worksheet(ws)
    assert cache[(4, 2)] == 60
    assert cache[(5, 2)] == 20
    assert cache[(6, 2)] == 3
    assert cache[(7, 2)] == 20


def test_division_and_zero_division():
    wb = Workbook()
    ws = wb.active

    ws["A1"] = 500
    ws["B1"] = 5
    ws["C1"] = "=A1/B1"

    ws["A2"] = 500
    ws["B2"] = 0
    ws["C2"] = "=A2/B2"

    cache = FormulaEvaluator.evaluate_worksheet(ws)
    assert cache[(1, 3)] == 100
    assert cache[(2, 3)] == "#DIV/0!"


def test_running_balance_with_opening_balance():
    wb = Workbook()
    ws = wb.active
    ws._opening_balance = 1000.0

    # Row 1: Header text (non-numeric, should resolve to opening_balance)
    ws["F1"] = "Net Value"
    ws["D2"] = 500
    ws["E2"] = 90
    ws["F2"] = "=F1+D2-E2"

    ws["D3"] = 300
    ws["E3"] = 54
    ws["F3"] = "=F2+D3-E3"

    cache = FormulaEvaluator.evaluate_worksheet(ws)
    assert cache[(2, 6)] == 1410
    assert cache[(3, 6)] == 1656


def test_complex_math_with_parentheses_and_power():
    wb = Workbook()
    ws = wb.active

    ws["A1"] = 100
    ws["B1"] = 20
    ws["C1"] = 2
    ws["D1"] = "=(A1 - B1) * (C1 ^ 3)"

    cache = FormulaEvaluator.evaluate_worksheet(ws)
    assert cache[(1, 4)] == 640


def test_sales_tax_template_formulas_fast_and_full():
    from app.engine.template_cache_manager import TemplateCacheManager
    from app.engine.template_processor import TemplateProcessor
    from app.engine.compiled_template import TemplateCompiler, CompiledRenderer
    import openpyxl

    wb = openpyxl.load_workbook("templates/SALES_TAX_INV.xlsx", data_only=False)

    data = {
        "opening_balance": 1000,
        "company": {"name": "HA MULTAN FOOD INDUSTRIES"},
        "invoices": [{
            "invoice_no": "INV-001",
            "items": [
                {"description": "Biscuits", "qty": 5, "item_value": 500, "st_value": 90},
                {"description": "Candy", "qty": 10, "item_value": 300, "st_value": 54},
                {"description": "Wafers", "qty": 4, "item_value": 200, "st_value": 36}
            ]
        }]
    }

    # 1. Full compiler mode via TemplateProcessor
    proc = TemplateProcessor(wb)
    out_wb = proc.process(data)
    out_ws = out_wb.active
    calc = FormulaEvaluator.evaluate_worksheet(out_ws)

    # Find the rows containing the item descriptions
    biscuits_row = None
    for r in range(1, out_ws.max_row + 1):
        if out_ws.cell(r, 1).value == "Biscuits":
            biscuits_row = r
            break
    assert biscuits_row is not None

    # Biscuits: Rate = 100 (500/5), Net Value = 1410 (1000 + 500 - 90)
    assert calc.get((biscuits_row, 3)) == 100
    assert calc.get((biscuits_row, 6)) == 1410

    # Candy: Rate = 30 (300/10), Net Value = 1656 (1410 + 300 - 54)
    assert calc.get((biscuits_row + 1, 3)) == 30
    assert calc.get((biscuits_row + 1, 6)) == 1656

    # Wafers: Rate = 50 (200/4), Net Value = 1820 (1656 + 200 - 36)
    assert calc.get((biscuits_row + 2, 3)) == 50
    assert calc.get((biscuits_row + 2, 6)) == 1820

    # Subtotal: Value = 1000, St Value = 180, Net Value = 4886
    assert calc.get((biscuits_row + 3, 4)) == 1000
    assert calc.get((biscuits_row + 3, 5)) == 180
    assert calc.get((biscuits_row + 3, 6)) == 4886

    # 2. Fast compiler mode via TemplateCompiler
    ast = TemplateCompiler.compile(wb, "TEST_SALES_TAX_VERIFY")
    html = CompiledRenderer.render(ast, data)

    assert "100" in html
    assert "1,410" in html
    assert "1,656" in html
    assert "1,820" in html
    assert "1,000" in html
    assert "180" in html
    assert "4,886" in html

