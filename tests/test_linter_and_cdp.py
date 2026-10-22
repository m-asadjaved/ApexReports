import openpyxl
import pytest
from fastapi.testclient import TestClient
from main import app
from app.engine.template_linter import TemplateLinter
from app.renderer.chrome_cdp_pool import ChromeCdpPool

client = TestClient(app)

def test_validate_invoice_report_endpoint():
    payload = {
        "template_code": "INVOICE_REPORT",
        "data": {
            "invoices": [
                {
                    "company": {"name": "Test Acme"},
                    "invoice_no": "INV-100",
                    "invoice_date": "2026-10-22",
                    "customer": {"name": "Test Customer"},
                    "items": [
                        {"item_code": "A1", "description": "Item 1", "qty": 1, "unit_price": 100}
                    ]
                }
            ]
        }
    }
    response = client.post("/api/v1/templates/validate", json=payload)
    assert response.status_code == 200
    data = response.json()
    assert data["valid"] is True
    assert data["score"] >= 90
    assert len(data["errors"]) == 0
    assert "summary" in data

def test_validate_nonexistent_template():
    payload = {
        "template_code": "UNKNOWN_TEMPLATE",
        "data": {}
    }
    response = client.post("/api/v1/templates/validate", json=payload)
    assert response.status_code == 404

def test_linter_detects_unclosed_block():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "{{#each items}}"
    ws["A2"] = "{{item_name}}"
    # Missing {{/each}}!
    
    result = TemplateLinter.lint_workbook(wb, template_code="TEST_BROKEN")
    assert result["valid"] is False
    assert result["score"] < 100
    assert any(e["code"] == "UNCLOSED_BLOCK" for e in result["errors"])

def test_linter_detects_mismatched_block_close():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "{{#if active}}"
    ws["A2"] = "Active Content"
    ws["A3"] = "{{/each}}"  # Mismatched close!
    
    result = TemplateLinter.lint_workbook(wb, template_code="TEST_MISMATCH")
    assert result["valid"] is False
    assert any(e["code"] == "MISMATCHED_BLOCK_CLOSE" for e in result["errors"])

def test_linter_detects_single_brace_typo():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "{customer_name}"  # Single brace typo!
    
    result = TemplateLinter.lint_workbook(wb, template_code="TEST_TYPO")
    assert any(w["code"] == "SINGLE_BRACE_TYPO" for w in result["warnings"])

def test_linter_detects_merged_cell_hazard():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Header"
    ws["A3"] = "{{#each items}}"
    ws["A4"] = "{{item_name}}"
    ws["A5"] = "{{/each}}"
    # Create dangerous vertical merge in column B crossing loop boundary (rows 2 to 6)
    ws.merge_cells("B2:B6")

    result = TemplateLinter.lint_workbook(wb, template_code="TEST_MERGE_HAZARD")
    assert any(w["code"] == "MERGED_CELL_LOOP_COLLISION" for w in result["warnings"])

def test_linter_detects_missing_array_in_payload():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "{{#each required_items}}"
    ws["A2"] = "{{desc}}"
    ws["A3"] = "{{/each}}"
    
    # Payload missing 'required_items'
    data = {"company": "Acme"}
    result = TemplateLinter.lint_workbook(wb, template_code="TEST_ARRAY", data=data)
    assert any(w["code"] == "MISSING_ARRAY_DATA" for w in result["warnings"])

def test_chrome_cdp_pool_availability_and_health():
    res = client.get("/api/v1/health")
    assert res.status_code == 200
    data = res.json()
    assert "pdf_backend" in data
    assert "cdp_pool_running" in data
    assert "cdp_port" in data
    assert data["cdp_port"] == 9222

def test_chrome_cdp_pool_render():
    if not ChromeCdpPool.is_available():
        pytest.skip("Chrome executable not found on host machine.")
        
    html = "<html><body style='font-family:sans-serif;'><h2>Fast CDP PDF Test</h2><p>Zero overhead rendering</p></body></html>"
    pdf_bytes = ChromeCdpPool.render_html_to_pdf(html)
    assert pdf_bytes.startswith(b"%PDF-")
    assert len(pdf_bytes) > 1000
