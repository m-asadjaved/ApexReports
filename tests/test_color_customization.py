import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.colors import Color
from app.renderer.fast_html_renderer import FastHtmlRenderer
from app.renderer.color_resolver import ExcelColorResolver

def test_direct_rgb_fill_and_font():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Custom Brand Banner"
    # Custom Gold background, White text
    ws["A1"].fill = PatternFill(fill_type="solid", fgColor="FFFFD700")
    ws["A1"].font = Font(color="FFFFFFFF", bold=True)
    
    html = FastHtmlRenderer.render(ws)
    assert "background-color: #FFD700" in html
    assert "color: #FFFFFF" in html

def test_theme_color_with_tint():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Theme Tint Cell"
    
    # Theme color 4 (Accent 1 Blue) with 40% lighter tint
    theme_color = Color(type="theme", theme=4, tint=0.4)
    ws["A1"].fill = PatternFill(fill_type="solid", fgColor=theme_color)
    
    html = FastHtmlRenderer.render(ws)
    # Ensure background-color is resolved to a hex color
    assert "background-color: #" in html
    # Should not be pure black or white
    assert "background-color: #000000" not in html

def test_indexed_color_resolution():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Indexed Red Cell"
    # Indexed 2 is Red
    ws["A1"].font = Font(color=Color(type="indexed", indexed=2))
    
    html = FastHtmlRenderer.render(ws)
    assert "color: #FF0000" in html

def test_unfilled_cell_has_no_background_color():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Transparent cell"
    
    html = FastHtmlRenderer.render(ws)
    # The cell must not have a forced background-color in its inline style
    td_line = [l for l in html.splitlines() if "Transparent cell" in l][0]
    assert "background-color" not in td_line

def test_color_resolver_unit_methods():
    # Direct RGB
    assert ExcelColorResolver.resolve_color(Color(rgb="FF123456")) == "#123456"
    # Black in openpyxl ARGB is 00000000 or FF000000
    assert ExcelColorResolver.resolve_color(Color(rgb="00000000")) == "#000000"
    # Unset pattern fill is transparent (None)
    assert ExcelColorResolver.resolve_fill(PatternFill()) is None
    # Darker tint (50% luminance of White is 808080)
    tinted_dark = ExcelColorResolver.apply_tint("FFFFFF", -0.5)
    assert tinted_dark == "808080"
    # Lighter tint (50% luminance of Black is 808080)
    tinted_light = ExcelColorResolver.apply_tint("000000", 0.5)
    assert tinted_light == "808080"
